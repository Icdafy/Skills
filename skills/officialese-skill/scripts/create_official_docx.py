#!/usr/bin/env python
"""Create a starter DOCX using the uploaded SOE official-document format.

The typography matches yiti-skill's ``references/format-rules.md`` (the shared
公文格式标准). Fonts are set in the same order as the manual Word workflow:
first each run's Chinese font by role (方正小标宋简体 / 黑体 / 楷体_GB2312 /
仿宋_GB2312), then round parentheses and their contents to 楷体_GB2312 at the size
of their position (二号 in titles, 三号 in text, 五号 in tables), then the footer
``-1-`` to 四号宋体, and finally Times New Roman for every Western slot of the body
so digits, Latin letters and ``%`` are Times New Roman while Chinese characters
keep their fonts. The footer is excluded from that last pass. Inline bold uses
``**...**``.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from embed_fonts import embed_bundled_fonts  # noqa: E402  随技能分发的字体嵌入器
import ensure_fonts  # noqa: E402  公文字体检测（缺失时提示运行 ensure_fonts.py 安装）


TITLE_FONT = "方正小标宋简体"
BODY_FONT = "仿宋_GB2312"
KAITI_FONT = "楷体_GB2312"
HEITI_FONT = "黑体"
SONGTI_FONT = "宋体"
# 西文字母、阿拉伯数字及 % 等半角符号统一 Times New Roman；中文走 eastAsia，
# Word 在同一 run 内按字符类型自动分派。页码 -1- 整体用四号宋体，不参与西文替换。
WESTERN_FONT = "Times New Roman"

TITLE_SIZE = 22  # 二号
BODY_SIZE = 16  # 三号
PAGE_NUMBER_SIZE = 14  # 四号
TABLE_SIZE = 10.5  # 五号：表格内括号外仿宋_GB2312，括号及括号内楷体_GB2312
BODY_LINE_PT = 28
TITLE_LINE_PT = 30  # 主标题、副标题/部门行及各级编号标题
TWO_CHAR_INDENT_PT = BODY_SIZE * 2
SIGNATURE_RIGHT_INDENT_PT = BODY_SIZE * 4  # 落款右空四字

STYLE_TITLE = "Official Title"
STYLE_SUBTITLE = "Official Subtitle"
STYLE_BODY = "Official Body"
STYLE_H1 = "Official Heading 1"
STYLE_H2 = "Official Heading 2"
STYLE_H3 = "Official Heading 3"
STYLE_H4 = "Official Heading 4"

BOLD_PATTERN = re.compile(r"\*\*(.+?)\*\*", re.S)
# 四级标题序号"（1）"随标题用仿宋_GB2312加粗，不按括号规则改楷体。
H4_SERIAL = re.compile(r"^[（(]\d+[）)]")
# 全角数字、字母和百分号转半角，保证最后一轮 Times New Roman 能作用到它们。
FULLWIDTH = {code: code - 0xFEE0 for code in range(0xFF10, 0xFF1A)}
FULLWIDTH.update({code: code - 0xFEE0 for code in range(0xFF21, 0xFF3B)})
FULLWIDTH.update({code: code - 0xFEE0 for code in range(0xFF41, 0xFF5B)})
FULLWIDTH[0xFF05] = ord("%")

# Times New Roman advance widths in em, for hanging-indent alignment.
TNR_EM = {".": 0.25, ",": 0.25, ":": 0.278, "-": 0.333, " ": 0.25, "/": 0.278, "%": 0.833}


def text_width_pt(text: str, size_pt: float = BODY_SIZE) -> float:
    """Width of a line set in Chinese font + Times New Roman at size_pt."""
    width = 0.0
    for char in text:
        if ord(char) < 128:
            width += (0.5 if char.isdigit() else TNR_EM.get(char, 0.5)) * size_pt
        else:
            width += size_pt
    return width


def set_run_font(run, east_asia_font: str, size_pt: float, bold: bool = False,
                 *, western_font: str = WESTERN_FONT, east_asia_hint: bool = True) -> None:
    """Set all font slots explicitly; the footer passes 宋体 for every slot."""
    run.font.name = western_font
    run.font.size = Pt(size_pt)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), western_font)
    rfonts.set(qn("w:hAnsi"), western_font)
    rfonts.set(qn("w:cs"), western_font)
    rfonts.set(qn("w:eastAsia"), east_asia_font)
    if east_asia_hint:
        # 中文引号、破折号、省略号按中文字体排；数字、字母、% 属 ASCII，仍用 Times New Roman。
        rfonts.set(qn("w:hint"), "eastAsia")
    sz_cs = rpr.find(qn("w:szCs"))
    if sz_cs is None:
        sz_cs = OxmlElement("w:szCs")
        rpr.append(sz_cs)
    sz_cs.set(qn("w:val"), str(int(size_pt * 2)))
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:eastAsia"), "zh-CN")


def set_style_font(style, font_name: str, size_pt: float, bold: bool = False) -> None:
    style.font.name = WESTERN_FONT
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for slot in ("ascii", "hAnsi", "cs"):
        rfonts.set(qn("w:" + slot), WESTERN_FONT)
    rfonts.set(qn("w:eastAsia"), font_name)


def parenthesis_mask(text: str) -> list[bool]:
    """True for every character inside a matched pair of round parentheses."""
    mask = [False] * len(text)
    stack = []
    for index, char in enumerate(text):
        if char in "（(":
            stack.append((index, char))
        elif char in "）)" and stack:
            if stack[-1][1] == {"）": "（", ")": "("}[char]:
                start, _ = stack.pop()
                mask[start:index + 1] = [True] * (index - start + 1)
    return mask


def add_formatted_text(paragraph, text: str, font: str, size: float,
                       bold: bool = False, paren_size: float | None = None,
                       *, keep_serial: bool = False) -> None:
    """成对中英文圆括号及其内容（含嵌套）统一楷体_GB2312，保留原文和加粗。

    括号片段字号为 ``paren_size``，默认与所在位置相同（标题二号、正文三号、表格
    五号）。未配对的括号保持原格式，不影响后续正文。``**...**`` 为行内加粗。
    ``keep_serial`` 用于四级标题：开头的"（1）"随标题字体。
    """
    paren_size = size if paren_size is None else paren_size
    text = str(text).translate(FULLWIDTH)
    segments = []
    pos = 0
    for match in BOLD_PATTERN.finditer(text):
        if match.start() > pos:
            segments.append((text[pos:match.start()], bold))
        segments.append((match.group(1), True))
        pos = match.end()
    if pos < len(text):
        segments.append((text[pos:], bold))

    # Match nested full-/half-width round parentheses across bold boundaries.
    plain = "".join(part for part, _ in segments)
    in_parentheses = parenthesis_mask(plain)
    serial = H4_SERIAL.match(plain) if keep_serial else None
    if serial:
        in_parentheses[:serial.end()] = [False] * serial.end()

    offset = 0
    for part, is_bold in segments:
        start = 0
        while start < len(part):
            special = in_parentheses[offset + start]
            end = start + 1
            while end < len(part) and in_parentheses[offset + end] == special:
                end += 1
            set_run_font(
                paragraph.add_run(part[start:end]),
                KAITI_FONT if special else font,
                paren_size if special else size,
                is_bold,
            )
            start = end
        offset += len(part)


def set_paragraph_format(
    paragraph,
    *,
    line_pt: int = BODY_LINE_PT,
    first_indent_pt: float = TWO_CHAR_INDENT_PT,
    alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
    left_indent_pt: float = 0,
    right_indent_pt: float = 0,
) -> None:
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(line_pt)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.first_line_indent = Pt(first_indent_pt)
    if left_indent_pt:
        fmt.left_indent = Pt(left_indent_pt)
    if right_indent_pt:
        fmt.right_indent = Pt(right_indent_pt)
    paragraph.alignment = alignment


def configure_style(
    doc: Document,
    name: str,
    font_name: str,
    size_pt: float,
    *,
    bold: bool = False,
    line_pt: int = BODY_LINE_PT,
    first_indent: bool = True,
    alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
) -> None:
    styles = doc.styles
    style = styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(style, font_name, size_pt, bold)
    fmt = style.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(line_pt)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.first_line_indent = Pt(TWO_CHAR_INDENT_PT) if first_indent else Pt(0)
    style.paragraph_format.alignment = alignment


def configure_styles(doc: Document) -> None:
    set_style_font(doc.styles["Normal"], BODY_FONT, BODY_SIZE, False)
    center = WD_ALIGN_PARAGRAPH.CENTER
    configure_style(doc, STYLE_TITLE, TITLE_FONT, TITLE_SIZE, line_pt=TITLE_LINE_PT,
                    first_indent=False, alignment=center)
    configure_style(doc, STYLE_SUBTITLE, KAITI_FONT, BODY_SIZE, line_pt=TITLE_LINE_PT,
                    first_indent=False, alignment=center)
    configure_style(doc, STYLE_BODY, BODY_FONT, BODY_SIZE)
    configure_style(doc, STYLE_H1, HEITI_FONT, BODY_SIZE, bold=False, line_pt=TITLE_LINE_PT)
    configure_style(doc, STYLE_H2, KAITI_FONT, BODY_SIZE, bold=True, line_pt=TITLE_LINE_PT)
    configure_style(doc, STYLE_H3, BODY_FONT, BODY_SIZE, bold=True, line_pt=TITLE_LINE_PT)
    configure_style(doc, STYLE_H4, BODY_FONT, BODY_SIZE, bold=True, line_pt=TITLE_LINE_PT)


def set_chinese_language(doc: Document) -> None:
    """Chinese typography defaults instead of python-docx's en-US/ja-JP template."""
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    lang = defaults.find(".//" + qn("w:lang")) if defaults is not None else None
    if lang is not None:
        lang.set(qn("w:eastAsia"), "zh-CN")
    theme_lang = doc.settings.element.find(qn("w:themeFontLang"))
    if theme_lang is not None:
        theme_lang.set(qn("w:eastAsia"), "zh-CN")


def add_text_paragraph(
    doc: Document,
    text: str,
    *,
    font: str = BODY_FONT,
    size: float = BODY_SIZE,
    bold: bool = False,
    style: str = STYLE_BODY,
    first_indent: bool = True,
    alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
    line_pt: int = BODY_LINE_PT,
    keep_serial: bool = False,
):
    p = doc.add_paragraph(style=style)
    set_paragraph_format(
        p,
        line_pt=line_pt,
        first_indent_pt=TWO_CHAR_INDENT_PT if first_indent else 0,
        alignment=alignment,
    )
    add_formatted_text(p, text, font, size, bold, keep_serial=keep_serial)
    return p


def add_center_line(doc: Document, text: str, font: str, size: float, style: str,
                    bold: bool = False) -> None:
    """标题区居中行；标题内括号与所在行同字号（主标题二号、副标题三号）。"""
    add_text_paragraph(doc, text, font=font, size=size, bold=bold, style=style,
                       first_indent=False, alignment=WD_ALIGN_PARAGRAPH.CENTER,
                       line_pt=TITLE_LINE_PT)


def add_title_lines(doc: Document, title: str, font: str = TITLE_FONT,
                    size: float = TITLE_SIZE, style: str = STYLE_TITLE) -> None:
    """标题可用换行符折行，每行一个居中段落。"""
    for line in str(title).splitlines() or [""]:
        add_center_line(doc, line, font, size, style)


def add_blank_line(doc: Document, line_pt: int = BODY_LINE_PT):
    p = doc.add_paragraph(style=STYLE_BODY)
    set_paragraph_format(p, line_pt=line_pt, first_indent_pt=0)
    return p


# CT_PPr children that follow w:autoSpaceDE/w:autoSpaceDN in schema order.
AUTO_SPACE_SUCCESSORS = (
    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
    "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)


def disable_auto_spacing(paragraph) -> None:
    """Stop Word adding CJK/Latin gaps so hanging indents align to the character."""
    ppr = paragraph._p.get_or_add_pPr()
    for tag, successors in (("w:autoSpaceDE", ("w:autoSpaceDN",) + AUTO_SPACE_SUCCESSORS),
                            ("w:autoSpaceDN", AUTO_SPACE_SUCCESSORS)):
        element = ppr.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            ppr.insert_element_before(element, *successors)
        element.set(qn("w:val"), "0")


def normalize_attachment_name(name: str) -> str:
    """附件名称不加书名号，后不加标点符号；顺带去掉调用方可能重复写入的
    「附件：」前缀和序号。"""
    cleaned = str(name).strip()
    serial = r"(?:\d+|[一二三四五六七八九十百零〇两]+)"
    # 仅识别明确的附件标签或带分隔符的序号，保留名称中的年份、数字。
    cleaned = re.sub(rf"^附件\s*(?:[:：]\s*|{serial}\s*[.、．:：]\s*)",
                     "", cleaned)
    cleaned = re.sub(rf"^(?:{serial}\s*[.、．:：]|[（(]{serial}[）)])\s*",
                     "", cleaned)
    cleaned = cleaned.replace("《", "").replace("》", "")
    return cleaned.strip().rstrip("。；;，,、.．：:！!？?… \t\r\n")


def add_page_field(paragraph) -> None:
    """整个 -1-（左右短横线、PAGE 域及显示数字）均为四号宋体，四个字体槽均为宋体。"""
    def songti(run):
        set_run_font(run, SONGTI_FONT, PAGE_NUMBER_SIZE, western_font=SONGTI_FONT,
                     east_asia_hint=False)

    songti(paragraph.add_run("-"))
    field_run = paragraph.add_run()
    songti(field_run)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    result = OxmlElement("w:t")
    result.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, result, end):
        field_run._r.append(element)
    songti(paragraph.add_run("-"))
    # Paragraph mark in 四号宋体 too, so the footer line height follows the page number.
    mark = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + slot), SONGTI_FONT)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), str(PAGE_NUMBER_SIZE * 2))
    mark.extend([fonts, size])
    paragraph._p.get_or_add_pPr().append(mark)


def setup_document(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3.7)
    section.bottom_margin = Cm(3.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.6)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(2.35)

    # 奇偶页不同：奇数页居右，偶数页居左。
    doc.settings.odd_and_even_pages_header_footer = True
    odd_footer = section.footer.paragraphs[0]
    odd_footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_field(odd_footer)
    even_footer = section.even_page_footer.paragraphs[0]
    even_footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_page_field(even_footer)


HEADINGS = {
    1: (HEITI_FONT, False, STYLE_H1),  # 一、黑体
    2: (KAITI_FONT, True, STYLE_H2),  # （一）楷体_GB2312加粗
    3: (BODY_FONT, True, STYLE_H3),  # 1.仿宋_GB2312加粗
    4: (BODY_FONT, True, STYLE_H4),  # （1）仿宋_GB2312加粗
}


def add_heading(doc: Document, text: str, level: int) -> None:
    font, bold, style = HEADINGS.get(level, HEADINGS[4])
    p = add_text_paragraph(doc, text, font=font, bold=bold, style=style,
                           line_pt=TITLE_LINE_PT, keep_serial=level >= 4)
    p.paragraph_format.keep_with_next = True  # 标题不单独留在页末


def column_widths_pt(all_rows: list, n_cols: int, total_pt: float) -> list[float]:
    """Content-based column widths that fill the text width, like Word's autofit.

    Short labels (序号、数字、"（工作日）"，不超过6个汉字宽) never wrap, so they set a
    column's minimum; the width left over is shared in proportion to how much
    longer each column's longest line is than its minimum.
    """
    padding = 12.0  # left + right cell margins
    short_limit = TABLE_SIZE * 6
    minimum, maximum = [], []
    for c_idx in range(n_cols):
        widths = [text_width_pt(line, TABLE_SIZE)
                  for row in all_rows if c_idx < len(row)
                  for line in str(row[c_idx]).replace("**", "").splitlines()]
        short = [w for w in widths if w <= short_limit]
        low = max(short + [TABLE_SIZE * 2]) + padding
        minimum.append(low)
        maximum.append(max(widths + [0.0]) + padding if widths else low)
        maximum[-1] = max(maximum[-1], low)
    if sum(maximum) <= total_pt:
        extra = total_pt - sum(maximum)
        return [w + extra * w / sum(maximum) for w in maximum]
    stretch = [hi - lo for lo, hi in zip(minimum, maximum)]
    free = total_pt - sum(minimum)
    if free <= 0 or not sum(stretch):
        return [total_pt * lo / sum(minimum) for lo in minimum]
    return [lo + free * s / sum(stretch) for lo, s in zip(minimum, stretch)]


def add_table(doc: Document, rows: list[list], header: bool = True) -> None:
    """表格内一律五号：括号外仿宋_GB2312，括号及括号内楷体_GB2312，数字、字母和 %
    走 Times New Roman。首行为表头：加粗并跨页重复；各行不跨页断开。单倍行距、
    无首行缩进、水平与垂直居中；表格铺满版心，列宽按内容分配。"""
    all_rows = [[("" if cell is None else str(cell)) for cell in row] for row in rows if row]
    if not all_rows:
        return
    if doc.paragraphs and doc.paragraphs[-1].text.strip():
        doc.paragraphs[-1].paragraph_format.keep_with_next = True  # 引导句与表格同页
    n_cols = max(len(row) for row in all_rows)
    table = doc.add_table(rows=len(all_rows), cols=n_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    # Fill the window width so the table never overflows the text area; columns
    # start from content-based widths and Word may still autofit them.
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "pct")
    tbl_w.set(qn("w:w"), "5000")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "autofit")
    tbl_pr.append(layout)
    section = doc.sections[-1]
    text_width = section.page_width.pt - section.left_margin.pt - section.right_margin.pt
    widths = column_widths_pt(all_rows, n_cols, text_width)
    for grid_col, width in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        grid_col.set(qn("w:w"), str(round(width * 20)))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Pt(width)

    for r_idx, row_data in enumerate(all_rows):
        is_header = header and r_idx == 0
        tr_pr = table.rows[r_idx]._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))  # 行不跨页断开
        if is_header:
            tr_pr.append(OxmlElement("w:tblHeader"))  # 跨页重复表头
        for c_idx in range(n_cols):
            cell = table.rows[r_idx].cells[c_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            text = row_data[c_idx] if c_idx < len(row_data) else ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fmt = p.paragraph_format
            fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
            fmt.space_before = Pt(0)
            fmt.space_after = Pt(0)
            fmt.first_line_indent = Pt(0)
            add_formatted_text(p, text, BODY_FONT, TABLE_SIZE, bold=is_header)


def add_block(doc: Document, item) -> None:
    """正文条目：字符串为正文段落；``{"table": [[...], ...]}`` 为表格（首行默认为
    表头，``"header": false`` 时不设表头）。"""
    if isinstance(item, dict) and "table" in item:
        add_table(doc, item["table"], header=bool(item.get("header", True)))
    else:
        add_text_paragraph(doc, str(item))


def add_sections(doc: Document, sections: list[dict]) -> None:
    for section in sections:
        title = section.get("heading") or section.get("title")
        if title:
            add_heading(doc, title, int(section.get("level", 1)))
        for para in section.get("paragraphs", []):
            add_block(doc, para)
        children = section.get("children", [])
        if children:
            add_sections(doc, children)


def add_attachment_line(doc: Document, text: str, text_start_pt: float, first_line_start_pt: float):
    """One attachment paragraph whose wrapped lines start at text_start_pt."""
    p = doc.add_paragraph(style=STYLE_BODY)
    set_paragraph_format(
        p,
        first_indent_pt=first_line_start_pt - text_start_pt,
        left_indent_pt=text_start_pt,
    )
    disable_auto_spacing(p)
    add_formatted_text(p, text, BODY_FONT, BODY_SIZE)
    return p


def add_attachments(doc: Document, attachments: list[str], *, keep_with_body: bool = False) -> None:
    """附件说明：正文下空一行，左空二字写"附件："。

    单份：附件：名称（回行与名称首字对齐）。
    多份：附件：1.名称 / 2.名称……，"2."与"1."对齐，每份回行与序号后的名称首字对齐。
    名称不加书名号，后不加标点符号。缩进按实际字宽（"附件："3字、"1."按
    Times New Roman 宽度）计算，并关闭中西文自动间距，保证对齐。
    """
    names = [name for a in attachments or [] if (name := normalize_attachment_name(a))]
    if not names:
        return
    if keep_with_body and doc.paragraphs:
        # 有落款时，正文末段随附件说明、落款同页，落款页不只剩附件和署名。
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
    add_blank_line(doc).paragraph_format.keep_with_next = keep_with_body
    label = "附件："
    label_start = TWO_CHAR_INDENT_PT
    number_start = label_start + text_width_pt(label)
    if len(names) == 1:
        add_attachment_line(doc, label + names[0], number_start, label_start)
        return
    lines = []
    for idx, name in enumerate(names, start=1):
        number = f"{idx}."
        text_start = number_start + text_width_pt(number)
        if idx == 1:
            lines.append(add_attachment_line(doc, f"{label}{number}{name}", text_start, label_start))
        else:
            lines.append(add_attachment_line(doc, f"{number}{name}", text_start, number_start))
    for line in lines[:-1]:
        line.paragraph_format.keep_with_next = True  # 多份附件说明不跨页拆开


def add_signature_block(doc: Document, issuer: str | None, date: str | None) -> None:
    """落款：正文或附件说明下空两行；署名右空四字，成文日期在署名下居中对齐。

    署名与日期都以两者中较宽者为基准居中：右缩进 = 4 字 + (较宽者 − 本行宽) / 2，
    宽度按中文字体 + Times New Roman 实际字宽计算，并关闭中西文自动间距。
    """
    issuer = str(issuer or "").strip()
    date = str(date or "").strip()
    if not issuer and not date:
        return
    # 落款不单独成页：附件说明（或正文末段）、两行空行与署名连在一起。
    if doc.paragraphs:
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
    for _ in range(2):
        add_blank_line(doc).paragraph_format.keep_with_next = True
    width = max(text_width_pt(issuer), text_width_pt(date))
    for text in (issuer, date):
        if not text:
            continue
        p = doc.add_paragraph(style=STYLE_BODY)
        set_paragraph_format(
            p,
            first_indent_pt=0,
            alignment=WD_ALIGN_PARAGRAPH.RIGHT,
            right_indent_pt=SIGNATURE_RIGHT_INDENT_PT + (width - text_width_pt(text)) / 2,
        )
        disable_auto_spacing(p)
        add_formatted_text(p, text, BODY_FONT, BODY_SIZE)
        p.paragraph_format.keep_with_next = text == issuer and bool(date)


def add_appendices(doc: Document, appendices: list[dict]) -> None:
    """附件正文另起一页：左上角顶格"附件："（多份时"附件1："），三号仿宋_GB2312；
    附件标题二号方正小标宋简体居中，可折行；标题下空一行后接正文。"""
    for idx, appendix in enumerate(appendices or [], start=1):
        default_label = "附件：" if len(appendices) == 1 else f"附件{idx}："
        label = add_text_paragraph(doc, appendix.get("label", default_label), first_indent=False)
        label.paragraph_format.page_break_before = True
        label.paragraph_format.keep_with_next = True
        add_title_lines(doc, appendix.get("title", ""))
        add_blank_line(doc, TITLE_LINE_PT)
        for para in appendix.get("body", []):
            add_block(doc, para)
        add_sections(doc, appendix.get("sections", []))


def apply_western_font_pass(doc: Document) -> None:
    """最后一轮：全文（含表格，不含页脚）西文槽位统一 Times New Roman，中文字体不变。
    与在 Word 中全选后把西文字体设为 Times New Roman 等效。"""
    for run in doc.element.body.iter(qn("w:r")):
        rfonts = run.get_or_add_rPr().get_or_add_rFonts()
        for slot in ("ascii", "hAnsi", "cs"):
            rfonts.set(qn("w:" + slot), WESTERN_FONT)


def build_docx(data: dict, output: Path) -> None:
    doc = Document()
    configure_styles(doc)
    set_chinese_language(doc)
    setup_document(doc)

    issuer_title = data.get("issuer_title")
    if issuer_title:
        add_title_lines(doc, issuer_title)

    add_title_lines(doc, data.get("title", "关于XXXX的通知"))

    subtitle = data.get("subtitle")
    if subtitle:
        add_title_lines(doc, subtitle, KAITI_FONT, BODY_SIZE, STYLE_SUBTITLE)

    add_blank_line(doc)

    recipient = data.get("recipient")
    if recipient:
        add_text_paragraph(doc, recipient, first_indent=False)

    for para in data.get("body", []):
        add_block(doc, para)

    add_sections(doc, data.get("sections", []))
    has_signature = bool(data.get("issuer") or data.get("date"))
    add_attachments(doc, data.get("attachments", []), keep_with_body=has_signature)
    add_signature_block(doc, data.get("issuer"), data.get("date"))
    add_appendices(doc, data.get("appendices") or [])
    # 最后一步：在各中文字体（方正小标宋简体、黑体、楷体_GB2312、仿宋_GB2312）
    # 与页脚四号宋体设置完成后，再对全文统一选择一次 Times New Roman。
    apply_western_font_pass(doc)

    doc.save(output)
    # 嵌入随附的可嵌入字体（仿宋_GB2312、楷体_GB2312），使交付件在未装这两款
    # 字体的机器上不掉字；方正小标宋许可禁止嵌入，自动跳过。校验不过则保留
    # 未嵌入版本，不影响生成。
    embed_bundled_fonts(output)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Create an SOE-style Chinese official DOCX starter.")
    parser.add_argument("output", type=Path, help="Output .docx path")
    parser.add_argument("--input", type=Path, help="JSON document specification")
    parser.add_argument("--title", help="Main title")
    parser.add_argument("--issuer-title", help="Issuing unit line above title")
    parser.add_argument("--subtitle", help="Centered subtitle or department line")
    parser.add_argument("--recipient", help="Recipient line, ending with Chinese colon")
    parser.add_argument("--body", action="append", default=[], help="Body paragraph; repeat as needed")
    parser.add_argument("--attachment", action="append", default=[], help="Attachment name; repeat as needed")
    parser.add_argument("--issuer", help="Signature unit")
    parser.add_argument("--date", help="Chinese date, e.g. 2026年6月8日")
    parser.add_argument("--doc-type", help="Optional document type note, e.g. 通知/请示/报告/函")
    parser.add_argument("--direction", choices=["上行文", "平行文", "下行文", "默认"], help="Optional writing direction note")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if args.input:
        data = json.loads(args.input.read_text(encoding="utf-8-sig"))
    else:
        data = {}

    for key in ["title", "issuer_title", "subtitle", "recipient", "issuer", "date", "doc_type", "direction"]:
        value = getattr(args, key)
        if value:
            data[key] = value
    if args.body:
        data["body"] = args.body
    if args.attachment:
        data["attachments"] = args.attachment
    build_docx(data, args.output)
    try:
        missing = [item["family"] for item in ensure_fonts.check()]
    except Exception:  # 检测失败不影响生成
        missing = []
    if missing:
        print(f"本机未安装：{'、'.join(missing)}，Word 打开会用替代字体显示；"
              f"运行 python {Path(__file__).with_name('ensure_fonts.py')} 从技能自带字体安装。",
              file=sys.stderr)


if __name__ == "__main__":
    main()
