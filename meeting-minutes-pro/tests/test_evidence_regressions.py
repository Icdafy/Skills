"""Behavioral regressions for omissions, stale evidence and publication boundaries."""
from __future__ import annotations

import contextlib
import io
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import types
import unittest
from unittest.mock import patch
from argparse import Namespace

SCRIPTS = Path(__file__).parents[1] / 'scripts'
sys.path.insert(0, str(SCRIPTS))
import audit_coverage as audit
import fact_check as facts
import refine_transcript as refine
import review_gate as gate
import run_state as state
import transcribe
import install_skill
import qa_reconcile as qa


class Files(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def write(self, name, text):
        p = self.root / name
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(text, encoding='utf-8')
        return p

    def json(self, name, value):
        return self.write(name, json.dumps(value, ensure_ascii=False))


class NumberTests(Files):
    def verify(self, source, body):
        a = self.write('source.txt', source)
        b = self.write('minutes.txt', '项目纪要\n　　一、业务情况\n　　' + body)
        with contextlib.redirect_stdout(io.StringIO()):
            return facts.verify(b, [a], [])

    def test_currency_mismatch(self):
        self.assertEqual(self.verify('价格100万元。', '价格100万美元。'), 1)

    def test_sign_mismatch(self):
        self.assertEqual(self.verify('增长-5%。', '增长5%。'), 1)

    def test_spoken_negative_keeps_sign(self):
        self.assertEqual(self.verify('毛利率负百分之五。', '毛利率-5%。'), 0)
        self.assertEqual(self.verify('毛利率负百分之五。', '毛利率5%。'), 1)
        self.assertEqual(self.verify('利润负五百万元。', '利润-500万元。'), 0)

    def test_percent_is_not_percentage_point(self):
        self.assertEqual(self.verify('提高5个百分点。', '提高5%。'), 1)

    def test_negative_percentage_points(self):
        self.assertEqual(self.verify('变动负五个百分点。', '变动-5个百分点。'), 0)
        self.assertEqual(self.verify('变动-5个百分点。', '变动5个百分点。'), 1)

    def test_cross_transcript_compare_checks_currency(self):
        a=self.write('a.txt','售价100万元。')
        b=self.write('b.txt','售价100万美元。')
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(facts.compare(a,b),1)

    def test_mantissa_does_not_cover_scaled_value(self):
        self.assertEqual(self.verify('员工3000人。', '营收3000万元。'), 1)

    def test_decimal_led_body_keeps_whole_value(self):
        self.assertEqual(self.verify('研发投入3.5亿元。', '3.5亿元用于研发。'), 0)

    def test_small_quantities_and_chinese_years(self):
        for source in ('客户3家，设备5台，周期2年。', '客户三家，设备五台，周期两年。', '二〇二五年成立。'):
            self.assertTrue(audit.Window(1, 'x', source).number_tokens, source)

    def test_units_and_scales_can_be_equivalent(self):
        self.assertEqual(self.verify('营收三千万。', '营收0.3亿元。'), 0)

    def test_point_wording_requires_resolution(self):
        self.assertEqual(self.verify('涨了三个点。', '涨了3%。'), 1)

    def test_whole_pipeline_rejects_original_counterexamples(self):
        pairs = [('毛利率30%，市占率30%。', '毛利率30%。'),
                 ('价格100万元。', '价格100万美元。'),
                 ('订单10台，售后100人。', '售后100人。'),
                 ('客户3家，设备5台，周期2年。', '讨论了业务情况。')]
        for source, body in pairs:
            transcript = self.write('source.txt', source)
            minutes = self.write('minutes.txt', '项目纪要\n　　一、业务情况\n　　' + body)
            windows, _ = audit.build_windows(transcript, 300, 1500)
            ledger = self.write('coverage.txt', f'窗口 1（{windows[0].label}）：纳入 总结\n')
            proc = subprocess.run([sys.executable, '-X', 'utf8', '-B', str(SCRIPTS/'check_all.py'),
                str(minutes), '--transcript', str(transcript), '--ledger', str(ledger)], capture_output=True)
            self.assertEqual(proc.returncode, 1, source)
            self.assertFalse(gate.read_json(self.root/'checks-summary.json')['all_passed'])


class ReviewPlanningTests(unittest.TestCase):
    def test_all_covers_missing_middle(self):
        stamps = [{'text': '开头', 'start': 0, 'end': 10}, {'text': '结尾', 'start': 110, 'end': 120}]
        clips = refine.plan_clips(stamps, [], select_all=True, pad=1.5, merge_gap=4, max_clip=30, duration=120)
        self.assertEqual(state.covered_seconds([(c.start,c.end) for c in clips]), 120)

    def test_all_covers_empty_primary(self):
        clips = refine.plan_clips([], [], select_all=True, pad=1, merge_gap=4, max_clip=30, duration=80)
        self.assertEqual(state.covered_seconds([(c.start,c.end) for c in clips]), 80)

    def test_union_does_not_double_count(self):
        self.assertEqual(state.covered_seconds([(0,10),(5,15),(20,30)]), 25)

    def test_all_rejects_budget_before_model_loading(self):
        proc = subprocess.run([sys.executable, '-X', 'utf8', '-B', str(SCRIPTS/'refine_transcript.py'),
            '--transcript', 'missing.json', '--source', 'missing.wav', '--output-dir', '.',
            '--all', '--budget-minutes', '1', '--dry-run'], capture_output=True)
        self.assertNotEqual(proc.returncode, 0)
        self.assertIn('--budget-minutes', proc.stderr.decode('utf-8'))

    def test_changed_qualifiers_and_units_flagged(self):
        for a,b in [('预计明年收入100万元。','去年已实现收入100万元。'),
                    ('售价100万元。','售价100万美元。'),
                    ('毛利率30%，市占率30%。','毛利率30%。')]:
            self.assertTrue(refine.compare_texts(a,b,[]))

    def test_meaningful_short_questions_survive(self):
        for text in ('多少钱？','盈利了吗？','能量产吗？','多少？'):
            self.assertTrue(qa.is_question(text))
        self.assertFalse(qa.is_question('是吧？'))


class CacheTests(Files):
    def test_content_and_config_bound_checkpoint(self):
        source = self.write('audio.wav', 'first-content')
        identity = state.fingerprint(source, {'model': 'a', 'context': '术语'})
        cache = self.root/'chunk.json'
        state.atomic_json(cache, {'identity': identity, 'start': 0, 'end': 1, 'text': '原稿'})
        self.assertIsNotNone(state.read_checkpoint(cache, identity, 0, 1))
        for config in ({'model':'b','context':'术语'}, {'model':'a','context':'新词'}):
            self.assertIsNone(state.read_checkpoint(cache, state.fingerprint(source,config),0,1))
        source.write_text('other-content',encoding='utf-8')
        self.assertIsNone(state.read_checkpoint(cache,state.fingerprint(source,{'model':'a','context':'术语'}),0,1))

    def test_corrupt_or_legacy_checkpoint_is_recomputed(self):
        cache=self.write('chunk.json','{')
        self.assertIsNone(state.read_checkpoint(cache,'new',0,1))
        state.atomic_json(cache, {'start':0,'end':1,'text':'legacy'})
        self.assertIsNone(state.read_checkpoint(cache,'new',0,1))

    def test_adapter_does_not_reuse_legacy_model_result(self):
        source=self.write('audio.wav','audio')
        cp=self.root/'chunks'; cp.mkdir()
        state.atomic_json(cp/'chunk_0001.json', {'engine':'qwen','start':0,'end':310,'text':'旧稿','stamps':[]})
        class Model:
            @classmethod
            def from_pretrained(cls,*a,**kw): return cls()
            def transcribe(self,**kw):
                return [types.SimpleNamespace(text='新稿',language='Chinese',time_stamps=[])]
        args=Namespace(device='cpu',model='new',clip_duration=None,chunk_seconds=None,max_new_tokens=None,
                       timestamps=True,aligner='new',language='Chinese',no_resume=False,context='新词',offline=False)
        with patch.dict(sys.modules,{'torch':types.ModuleType('torch'),'qwen_asr':types.SimpleNamespace(Qwen3ASRModel=Model)}), \
             patch.object(transcribe,'choose_device',return_value=('cpu','float32')), \
             patch.object(transcribe,'detect_silences',return_value=[]), \
             patch.object(transcribe,'plan_chunks',return_value=[(0,310),(310,620)]), \
             patch.object(transcribe,'prepare_wav'):
            out=transcribe.run_qwen(args,source,620,cp)
        self.assertEqual(out.text,'新稿\n新稿')

    def test_offline_guard_blocks_network_in_child(self):
        code=f"import sys;sys.path.insert(0,{str(SCRIPTS)!r});import run_state,socket;run_state.enable_offline();socket.create_connection(('example.com',443))"
        proc=subprocess.run([sys.executable,'-X','utf8','-B','-c',code],capture_output=True)
        self.assertNotEqual(proc.returncode,0)
        self.assertIn(b'offline mode',proc.stderr)


class GateTests(Files):
    def test_missing_ledger_cannot_pass(self):
        source=self.write('source.txt','讨论了生产安排。')
        minutes=self.write('minutes.txt','生产会议纪要\n　　一、生产安排\n　　讨论了生产安排。')
        proc=subprocess.run([sys.executable,'-X','utf8','-B',str(SCRIPTS/'check_all.py'),str(minutes),
                             '--transcript',str(source)],capture_output=True)
        self.assertEqual(proc.returncode,1)
        self.assertFalse(gate.read_json(self.root/'checks-summary.json')['all_passed'])

    def test_split_transcript_evidence_is_rejected(self):
        source=self.json('source.json',{'text':'收入100万元。','timestamps':[{'text':'收入10万元。'}]})
        txt=self.write('source.txt','收入20万元。')
        self.assertEqual(len(gate.transcript_consistency(source,txt)),2)

    def test_pending_review_and_stale_bindings_fail(self):
        source=self.write('source.txt','毛利率30%，市占率30%。')
        minutes=self.write('minutes.txt','毛利率30%，市占率30%。')
        expected=gate.build_review(gate.input_bindings({'minutes':minutes}),{},gate.fact_rows(source),[])
        self.assertEqual(len(expected['facts']),2)
        self.assertTrue(gate.validate_review(expected,expected,minutes))
        review=json.loads(json.dumps(expected))
        review['inputs']={}
        self.assertIn('版本',gate.validate_review(review,expected,minutes)[0])

    def test_each_occurrence_must_point_to_real_text(self):
        source=self.write('source.txt','毛利率30%，市占率30%。')
        minutes=self.write('minutes.txt','毛利率30%，市占率30%。')
        expected=gate.build_review({}, {},gate.fact_rows(source),[])
        review=json.loads(json.dumps(expected))
        for row in review['facts']:
            row.update(decision='include',minutes_lines=[1],reason='已逐项核对对象、单位、时间及限定条件',reviewer='test')
        review['document_review']={'confirmed':True,'reason':'已核对全文与版式','reviewer':'test'}
        self.assertEqual(gate.validate_review(review,expected,minutes),[])
        review['facts'][1]['minutes_lines']=[100]
        self.assertTrue(gate.validate_review(review,expected,minutes))

    def test_audio_finding_needs_actual_listening_method(self):
        minutes=self.write('minutes.txt','会议纪要')
        expected=gate.build_review({}, {},[],[gate.finding('audio','分歧',audio=True)])
        review=json.loads(json.dumps(expected))
        review['findings'][0].update(decision='confirmed',reason='确认',reviewer='test',method='asr_again')
        review['document_review']={'confirmed':True,'reason':'已核对','reviewer':'test'}
        self.assertTrue(gate.validate_review(review,expected,minutes))

    def test_partial_high_risk_report_is_rejected(self):
        audio=self.write('source.wav','audio')
        source=self.json('source.json',{'source_sha256':state.file_hash(audio),'duration_seconds':120,'engine':'funasr'})
        report=self.json('refine.json',{'source_sha256':state.file_hash(audio),'transcript_sha256':state.file_hash(source),
            'primary_engine':'funasr','model':'Qwen','review_mode':'all','duration_seconds':120,'clips':[
                {'start':0,'end':10,'status':'consistent'},{'start':110,'end':120,'status':'consistent'}]})
        errors,_=gate.audio_evidence(audio,source,None,report,None,'high')
        self.assertTrue(any('完整录音' in e for e in errors))

    def test_empty_model_and_text_evidence_cannot_claim_full_review(self):
        audio=self.write('source.wav','audio')
        source=self.json('source.json',{'source_sha256':state.file_hash(audio),'duration_seconds':120,'engine':'funasr','model':'paraformer-zh'})
        report=self.json('refine.json',{'source_sha256':state.file_hash(audio),'transcript_sha256':state.file_hash(source),
            'primary_engine':'funasr','review_mode':'all','duration_seconds':120,
            'clips':[{'start':0,'end':120,'status':'consistent'}]})
        errors,_=gate.audio_evidence(audio,source,None,report,None,'high')
        self.assertTrue(any('模型' in e for e in errors))
        self.assertTrue(any('两引擎文本' in e for e in errors))

    def test_another_asr_pass_is_not_a_listened_revision(self):
        audio=self.write('source.wav','audio')
        base={'source_sha256':state.file_hash(audio),'duration_seconds':10,'engine':'funasr','model':'paraformer-zh'}
        raw=self.json('raw.json',{**base,'text':'原稿'})
        accepted=self.json('accepted.json',{**base,'text':'修订稿'})
        revisions=self.json('revisions.json',{'raw_sha256':state.file_hash(raw),'accepted_sha256':state.file_hash(accepted),
            'entries':[{'start':0,'end':5,'before':'原稿','after':'修订稿','reviewer':'test','reason':'再次识别','method':'asr_again'}]})
        errors,_=gate.audio_evidence(audio,accepted,raw,None,revisions,'standard')
        self.assertTrue(any('实际回听' in e for e in errors))

    def test_non_rendered_docx_cannot_release(self):
        self.assertTrue(gate.validate_render(None,None))

    def test_unused_heading_font_is_not_required_in_pdf(self):
        from docx import Document
        from docx.oxml.ns import qn
        from render_docx import used_font_markers
        document=Document()
        run=document.add_paragraph().add_run('正文')
        run.font.name='仿宋_GB2312'
        run._element.rPr.rFonts.set(qn('w:eastAsia'),'仿宋_GB2312')
        path=self.root/'body.docx';document.save(path)
        self.assertEqual(used_font_markers(path),('FangSong_GB2312',))


class InstallTests(Files):
    def test_update_preserves_private_data_and_does_not_copy_other_projects(self):
        source=self.root/'source'; target=self.root/'target'
        self.write('source/SKILL.md','new')
        self.write('source/glossary/private.txt','source secret')
        self.write('source/glossary/banned-phrases.txt','source policy')
        self.write('source/glossary/industry/public.txt','public')
        self.write('target/SKILL.md','old')
        private=self.write('target/glossary/own.txt','target secret')
        banned=self.write('target/glossary/banned-phrases.txt','target policy')
        install_skill.install_to(source,target,force=True)
        self.assertEqual(private.read_text(),'target secret')
        self.assertEqual(banned.read_text(),'target policy')
        self.assertFalse((target/'glossary/private.txt').exists())
        self.assertTrue((target/'glossary/industry/public.txt').exists())
        self.assertEqual((target/'SKILL.md').read_text(),'new')

    def test_nested_install_is_rejected(self):
        source=self.root/'source';source.mkdir()
        with self.assertRaises(ValueError):
            install_skill.install_to(source,source/'child',True)


if __name__ == '__main__':
    unittest.main()
