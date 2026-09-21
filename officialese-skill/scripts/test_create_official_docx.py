"""Regression checks for typography in the saved DOCX (stdlib unittest)."""

import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn

import create_official_docx as generator


class OfficialFormatTests(unittest.TestCase):
    def assert_font(self, run, font, size):
        fonts = run._r.rPr.rFonts
        for script in ('ascii', 'hAnsi', 'cs', 'eastAsia'):
            self.assertEqual(fonts.get(qn('w:' + script)), font)
        self.assertEqual(run.font.size.pt, size)

    def assert_mixed_font(self, run, east_asia, size):
        """中文用 ``east_asia``，数字、字母、% 等用 Times New Roman。"""
        fonts = run._r.rPr.rFonts
        self.assertEqual(fonts.get(qn('w:eastAsia')), east_asia)
        for script in ('ascii', 'hAnsi', 'cs'):
            self.assertEqual(fonts.get(qn('w:' + script)), generator.WESTERN_FONT)
        self.assertEqual(run.font.size.pt, size)

    def test_parentheses_preserve_text_and_surrounding_format(self):
        doc = Document()
        p = doc.add_paragraph()
        text = '正文ABC（外层(内层123)说明）后文(ABC 2)末尾（未闭合'
        generator.add_formatted_text(p, text, generator.BODY_FONT, 16, True)
        self.assertEqual(p.text, text)
        for run in p.runs:
            if run.text.startswith(('（外层', '(ABC')):
                self.assert_mixed_font(run, generator.KAITI_FONT, 16)
            else:
                self.assertEqual(run._r.rPr.rFonts.get(qn('w:eastAsia')),
                                 generator.BODY_FONT)
                self.assertEqual(run._r.rPr.rFonts.get(qn('w:ascii')),
                                 generator.WESTERN_FONT)
            self.assertTrue(run.bold)

    def test_attachment_names_and_counts(self):
        cases = {
            '附件1.《实施方案（试行）》。': '实施方案（试行）',
            '附件：一、申报表；': '申报表',
            '附件二：测算表！': '测算表',
            '（三）《2026年度计划》？': '2026年度计划',
            '2026年度计划': '2026年度计划',
            '关于《实施方案》的说明：': '关于实施方案的说明',
            '附件2026年度说明': '附件2026年度说明',
        }
        for source, expected in cases.items():
            with self.subTest(source=source):
                self.assertEqual(generator.normalize_attachment_name(source), expected)
        doc = Document()
        generator.configure_styles(doc)
        generator.add_attachments(doc, ['附件1.《实施方案（试行）》。', '《》'])
        self.assertEqual([p.text for p in doc.paragraphs if p.text],
                         ['附件：实施方案（试行）'])
        doc = Document()
        generator.configure_styles(doc)
        generator.add_attachments(doc, ['附件一：申报表；', '附件2.测算表！'])
        items = [p for p in doc.paragraphs if p.text]
        self.assertEqual([p.text for p in items], ['附件：1.申报表', '2.测算表'])
        self.assertEqual(items[0].paragraph_format.left_indent.pt, 96)
        self.assertEqual(items[1].paragraph_format.first_line_indent.pt, -16)

    def test_saved_document_spacing_parentheses_and_whole_page_field(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'format.docx'
            generator.build_docx({
                'title': '关于工作安排的通知（试行2026）',
                'issuer_title': '公司文件',
                'subtitle': '业务部门',
                'recipient': '各部门（含子公司）：',
                'body': ['正文（说明ABC 123）结束'],
                'sections': [{'heading': '（1）工作要求', 'level': 4,
                              'paragraphs': [{'table': [['项目', '占比（%）'],
                                                        ['研发投入（2026年）', '12.5%']]}]}],
                'attachments': ['附件1.《工作计划（试行）》；'],
                'issuer': '某某公司（集团）',
                'date': '2026年9月14日',
            }, path)
            doc = Document(path)
            parenthetical_runs = []
            for paragraph in doc.paragraphs:
                expected_spacing = 30 if paragraph.style.name in (
                    generator.STYLE_TITLE, generator.STYLE_SUBTITLE) else 28
                self.assertEqual(paragraph.paragraph_format.line_spacing.pt,
                                 expected_spacing)
                self.assertEqual(paragraph.paragraph_format.line_spacing_rule,
                                 WD_LINE_SPACING.EXACTLY)
                for run in paragraph.runs:
                    if run.text.startswith(('（', '(')):
                        parenthetical_runs.append(run.text)
                        self.assert_mixed_font(run, generator.KAITI_FONT, 16)
                    else:
                        self.assertEqual(run._r.rPr.rFonts.get(qn('w:ascii')),
                                         generator.WESTERN_FONT)
            self.assertEqual(len(parenthetical_runs), 6)
            heading4 = next(p for p in doc.paragraphs if p.text == '（1）工作要求')
            self.assertTrue(all(run.bold for run in heading4.runs))
            self.assertEqual(heading4.runs[-1]._r.rPr.rFonts.get(qn('w:eastAsia')),
                             generator.BODY_FONT)
            cells = [p for row in doc.tables[0].rows for c in row.cells
                     for p in c.paragraphs]
            self.assertEqual([p.text for p in cells],
                             ['项目', '占比（%）', '研发投入（2026年）', '12.5%'])
            for paragraph in cells:
                for run in paragraph.runs:
                    east = (generator.KAITI_FONT if run.text.startswith('（')
                            else generator.BODY_FONT)
                    self.assert_mixed_font(run, east, 10.5)
                    self.assertFalse(run.bold)
            title = next(p for p in doc.paragraphs if p.text.startswith('关于'))
            self.assertEqual(title.runs[0].font.size.pt, 22)
            section = doc.sections[0]
            for footer in (section.footer, section.even_page_footer):
                paragraph = footer.paragraphs[0]
                self.assertEqual(paragraph.text, '-1-')
                self.assertEqual(len(paragraph.runs), 3)
                for run in paragraph.runs:
                    self.assert_font(run, generator.SONGTI_FONT, 14)
                instructions = paragraph._p.xpath('.//w:instrText')
                self.assertEqual([node.text.strip() for node in instructions], ['PAGE'])


if __name__ == '__main__':
    unittest.main()
