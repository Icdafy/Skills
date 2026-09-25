"""Validate yiti's generated DOCX font slots, spacing, attachments, signature, fields and text checks."""
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from docx import Document
from docx.oxml.ns import qn

REPO = Path(__file__).resolve().parents[2]
SKILL = REPO / 'yiti-skill'
GENERATOR = SKILL / 'scripts/create_yiti_docx.py'
CHECKER = SKILL / 'scripts/check_yiti_text.py'
EXIT_SPEC = SKILL / 'assets/examples/exit-yiti-spec.json'


def run_generator(spec: Path, out: Path) -> None:
    proc = subprocess.run(
        [sys.executable, str(GENERATOR), str(spec), str(out)],
        capture_output=True, text=True, encoding="utf-8", errors="replace",
    )
    if proc.returncode:
        raise AssertionError(proc.stdout + proc.stderr)


def twips(paragraph, attr):
    ind = paragraph._p.pPr.find(qn('w:ind'))
    value = ind.get(qn('w:' + attr)) if ind is not None else None
    return int(value) if value is not None else 0


class YitiFormatTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory(prefix="yiti-format-", dir=REPO.parent)
        cls.addClassCleanup(cls.tmp.cleanup)
        data = {
            "title": "审议关于示例公司（有限合伙ABC123）的议题",
            "intro": "正文2026（含**加粗ABC12**及(嵌套34)内容）恢复正文56。",
            "blocks": [
                {"type": "h1", "text": "一、基本信息"},
                {"type": "h2", "text": "（一）议案"},
                {"type": "h3", "text": "1.依据"},
                {"type": "h4", "text": "（1）条款（修订）"},
                {"type": "para", "text": "英文括号(ABC123)后文。未闭合（保留原文，全角２０２６％"},
                {"type": "table", "caption": "指标（单位元）：",
                 "header": ["名称", "金额"], "rows": [["利润（万元123）", "123.45"]]},
            ],
        }
        cls.docs = []
        for index, names in enumerate((
            ["《通知》。", "关于《议案》的说明；", "表决票（修订版ABC2）。"],
            ["《会议通知》。"], [],
        )):
            spec = Path(cls.tmp.name) / f"input{index}.json"
            out = spec.with_suffix('.docx')
            spec.write_text(json.dumps({**data, "attachments": names}, ensure_ascii=False), encoding="utf-8")
            run_generator(spec, out)
            cls.docs.append(Document(out))
        cls.exit_out = Path(cls.tmp.name) / 'exit.docx'
        run_generator(EXIT_SPEC, cls.exit_out)
        cls.exit_doc = Document(cls.exit_out)

    def assert_font(self, run, east_asia, western, size, bold=None):
        fonts = run._r.rPr.rFonts
        for slot in ('ascii', 'hAnsi', 'cs'):
            self.assertEqual(fonts.get(qn('w:' + slot)), western, run.text)
        self.assertEqual(fonts.get(qn('w:eastAsia')), east_asia, run.text)
        self.assertEqual(run.font.size.pt, size, run.text)
        self.assertEqual(run._r.rPr.find(qn('w:szCs')).get(qn('w:val')), str(int(size * 2)))
        if bold is not None:
            self.assertEqual(bool(run.bold), bold, run.text)

    def test_parentheses_kaiti_with_times_new_roman_digits(self):
        doc = self.docs[0]
        body = [(p, 16) for p in doc.paragraphs]
        body += [(p, 10.5) for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs]
        checked = 0
        for paragraph, size in body:
            # Test every run against the visible paired span, independently
            # of how the generator chose to split the runs.
            text = paragraph.text
            opener = next((i for i, ch in enumerate(text) if ch in '（('), -1)
            closer = max(text.rfind('）'), text.rfind(')'))
            if text.startswith('（1）'):
                opener = text.find('（', 3)
            offset = 0
            for run in paragraph.runs:
                if opener >= 0 and closer > opener and offset <= closer and offset + len(run.text) > opener:
                    self.assert_font(run, '楷体_GB2312', 'Times New Roman', size)
                    checked += 1
                offset += len(run.text)
        self.assertGreater(checked, 5)
        intro = next(p for p in doc.paragraphs if p.text.startswith('正文2026'))
        self.assertEqual(intro.text, '正文2026（含加粗ABC12及(嵌套34)内容）恢复正文56。')
        self.assertTrue(next(r for r in intro.runs if r.text == '加粗ABC12').bold)
        self.assert_font(intro.runs[-1], '仿宋_GB2312', 'Times New Roman', 16)
        self.assert_font(doc.paragraphs[0].runs[0], '方正小标宋简体', 'Times New Roman', 22)

    def test_table_text_is_wuhao_fangsong_and_kaiti(self):
        table = self.docs[0].tables[0]
        cell_runs = table.cell(1, 0).paragraphs[0].runs
        self.assertEqual([r.text for r in cell_runs], ['利润', '（万元123）'])
        self.assert_font(cell_runs[0], '仿宋_GB2312', 'Times New Roman', 10.5)
        self.assert_font(cell_runs[1], '楷体_GB2312', 'Times New Roman', 10.5)
        self.assert_font(table.cell(1, 1).paragraphs[0].runs[0], '仿宋_GB2312', 'Times New Roman', 10.5)
        self.assert_font(table.cell(0, 0).paragraphs[0].runs[0], '仿宋_GB2312', 'Times New Roman', 10.5, bold=True)
        tr_pr = table.rows[0]._tr.trPr
        self.assertIsNotNone(tr_pr.find(qn('w:tblHeader')))

    def test_heading_fonts_including_level_four(self):
        paragraphs = {p.text: p for p in self.docs[0].paragraphs}
        self.assert_font(paragraphs['一、基本信息'].runs[0], '黑体', 'Times New Roman', 16, bold=False)
        self.assert_font(paragraphs['（一）议案'].runs[0], '楷体_GB2312', 'Times New Roman', 16, bold=True)
        self.assert_font(paragraphs['1.依据'].runs[0], '仿宋_GB2312', 'Times New Roman', 16, bold=True)
        h4 = paragraphs['（1）条款（修订）']
        self.assertEqual(h4.runs[0].text, '（1）条款')
        self.assert_font(h4.runs[0], '仿宋_GB2312', 'Times New Roman', 16, bold=True)
        self.assert_font(h4.runs[1], '楷体_GB2312', 'Times New Roman', 16, bold=True)
        for text in ('一、基本信息', '（一）议案', '1.依据', '（1）条款（修订）'):
            self.assertTrue(paragraphs[text].paragraph_format.keep_with_next, text)

    def test_times_new_roman_pass_and_fullwidth_digits(self):
        for doc in (self.docs[0], self.exit_doc):
            for run in doc.element.body.iter(qn('w:r')):
                fonts = run.rPr.rFonts
                for slot in ('ascii', 'hAnsi', 'cs'):
                    self.assertEqual(fonts.get(qn('w:' + slot)), 'Times New Roman')
        para = next(p for p in self.docs[0].paragraphs if p.text.startswith('英文括号'))
        self.assertTrue(para.text.endswith('全角2026%'))

    def test_fixed_title_and_body_spacing(self):
        heading_texts = {'一、基本信息', '（一）议案', '1.依据', '（1）条款（修订）'}
        for index, paragraph in enumerate(self.docs[0].paragraphs):
            spacing = paragraph._p.pPr.find(qn('w:spacing'))
            expected = 600 if index < 2 or paragraph.text in heading_texts else 560
            self.assertEqual(spacing.get(qn('w:lineRule')), 'exact')
            self.assertEqual(spacing.get(qn('w:line')), str(expected), paragraph.text)

    def test_attachment_labels_names_and_alignment(self):
        expected = [
            ['附件：1.通知', '2.关于议案的说明', '3.表决票（修订版ABC2）'],
            ['附件：会议通知'], [],
        ]
        for doc, names in zip(self.docs, expected):
            texts = [p.text for p in doc.paragraphs]
            start = next((i for i, t in enumerate(texts) if t.startswith('附件')), len(texts))
            self.assertEqual(texts[start:], names)
            if names:
                self.assertEqual(texts[start - 1], '')  # 正文下空一行
        multi = self.docs[0].paragraphs[-3:]
        # "附件：" starts at 2 chars (640), occupies 3 chars (960); "1." is 12 pt (240) in Times New Roman.
        self.assertEqual((twips(multi[0], 'left'), twips(multi[0], 'hanging')), (1840, 1200))
        for paragraph in multi[1:]:
            self.assertEqual((twips(paragraph, 'left'), twips(paragraph, 'hanging')), (1840, 240))
        single = self.docs[1].paragraphs[-1]
        self.assertEqual((twips(single, 'left'), twips(single, 'hanging')), (1600, 960))
        for paragraph in multi:
            self.assert_font(paragraph.runs[0], '仿宋_GB2312', 'Times New Roman', 16)

    def test_complete_footer_font_and_dynamic_field(self):
        for doc in (self.docs[0], self.exit_doc):
            self.assertTrue(doc.settings.odd_and_even_pages_header_footer)
            section = doc.sections[0]
            for footer, align in ((section.footer, 'right'), (section.even_page_footer, 'left')):
                p = footer.paragraphs[0]
                self.assertEqual(p.text, '-1-')
                self.assertEqual(p._p.pPr.find(qn('w:jc')).get(qn('w:val')), align)
                for run in p.runs:
                    self.assert_font(run, '宋体', '宋体', 14)
                self.assertEqual([e.text.strip() for e in p._p.iter(qn('w:instrText'))], ['PAGE'])
                self.assertEqual([e.get(qn('w:fldCharType')) for e in p._p.iter(qn('w:fldChar'))],
                                 ['begin', 'separate', 'end'])

    def test_exit_kind_structure_signature_and_appendix(self):
        texts = [p.text for p in self.exit_doc.paragraphs]
        self.assertEqual(texts[:4], ['审议关于XX智能科技有限公司', '投资项目退出的议题', '提交子公司：投管公司', ''])
        self.assertNotIn('投委会：', texts)
        self.assertTrue(texts[4].startswith('2024年X月'))
        self.assertTrue(texts[5].endswith('现提请投委会审议。具体如下：'))
        attach = texts.index('附件：XX智能科技有限公司项目退出方案')
        self.assertEqual(texts[attach + 1:attach + 5], ['', '', 'XX投资管理有限公司', '2026年9月X日'])
        issuer, date = self.exit_doc.paragraphs[attach + 3], self.exit_doc.paragraphs[attach + 4]
        self.assertEqual(twips(issuer, 'right'), 1280)  # 署名右空四字
        self.assertGreater(twips(date, 'right'), twips(issuer, 'right'))  # 日期居中于署名下
        for index in range(attach - 2, attach + 4):  # 正文末段、附件说明与落款同页
            self.assertTrue(self.exit_doc.paragraphs[index].paragraph_format.keep_with_next, texts[index])
        label = self.exit_doc.paragraphs[attach + 5]
        self.assertEqual(label.text, '附件：')
        self.assertTrue(label.paragraph_format.page_break_before)
        self.assertEqual(texts[attach + 6:attach + 8], ['XX智能科技有限公司项目', '退出方案'])
        self.assert_font(self.exit_doc.paragraphs[attach + 6].runs[0], '方正小标宋简体', 'Times New Roman', 22)
        table = self.exit_doc.tables[0]
        widths = [int(g.get(qn('w:w'))) for g in table._tbl.tblGrid.findall(qn('w:gridCol'))]
        self.assertLess(widths[0], widths[1])  # 序号窄、事项宽
        self.assertLessEqual(abs(sum(widths) - 8844), 4)  # 铺满版心不溢出


class YitiTextCheckTests(unittest.TestCase):
    def run_checker(self, text: str, suffix='.md'):
        with tempfile.TemporaryDirectory(dir=REPO.parent) as tmp:
            path = Path(tmp) / f'draft{suffix}'
            path.write_text(text, encoding='utf-8')
            return subprocess.run([sys.executable, str(CHECKER), str(path)], capture_output=True,
                                  text=True, encoding='utf-8', errors='replace')

    def test_deferral_endings_and_contrast_patterns_are_hard_hits(self):
        for sentence in ('公司经营承压，具体影响需要后续进行判断。',
                         '相关风险有待进一步研判。',
                         '后续需持续关注回款情况。',
                         '退出时机视情况而定。',
                         '本次退出不是被动止损，而是主动防范风险。',
                         '采用现金回购而非股权置换。',
                         '值得注意的是，公司现金流偏紧。'):
            with self.subTest(sentence=sentence):
                proc = self.run_checker(sentence)
                self.assertEqual(proc.returncode, 1, proc.stdout)
                self.assertIn('[硬规则]', proc.stdout)

    def test_rule_statements_and_final_style_pass(self):
        text = ('最终价格以审计、评估结论及国资备案结果为准。\n'
                '收入转化和现金流贡献尚需较长时间。\n'
                '判断公司全年业绩目标基本已无实现可能，届时将大概率触发平价回购条款。\n')
        proc = self.run_checker(text)
        self.assertEqual(proc.returncode, 0, proc.stdout)
        self.assertNotIn('[警告]', proc.stdout)

    def test_tone_warnings(self):
        proc = self.run_checker('项目必然失败，公司或许会继续亏损。')
        self.assertEqual(proc.returncode, 0, proc.stdout)
        self.assertIn('语气偏激进', proc.stdout)
        self.assertIn('语气偏保守', proc.stdout)

    def test_bundled_exit_example_is_clean(self):
        proc = subprocess.run([sys.executable, str(CHECKER), str(EXIT_SPEC)], capture_output=True,
                              text=True, encoding='utf-8', errors='replace')
        self.assertEqual(proc.returncode, 0, proc.stdout)
        self.assertIn('硬规则命中 0 处，警告 0 处', proc.stdout)


if __name__ == '__main__':
    unittest.main()
