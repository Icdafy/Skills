"""Saved-document regression tests for the three investment-report skills."""
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from docx import Document
from docx.enum.text import WD_LINE_SPACING, WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

REPO = Path(__file__).resolve().parents[2]
SKILLS = ('gongsi-qingkuang', 'hangye-fenxi', 'zhuying-yewu-fenxi')


class PublicDocumentFormatTests(unittest.TestCase):
    def assert_font(self, run, font, size):
        for attribute in ('ascii', 'hAnsi', 'cs'):
            self.assertEqual(run._r.rPr.rFonts.get(qn('w:' + attribute)), 'Times New Roman')
        self.assertEqual(run._r.rPr.rFonts.get(qn('w:eastAsia')), font)
        self.assertEqual(run.font.size.pt, size)

    def test_final_western_pass_preserves_chinese_and_field_format(self):
        spec = importlib.util.spec_from_file_location(
            'report_builder', REPO / SKILLS[0] / 'scripts/build_docx.py')
        builder = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(builder)
        doc = Document()
        places = [doc.add_paragraph(), doc.sections[0].header.paragraphs[0],
                  doc.sections[0].footer.paragraphs[0],
                  doc.add_table(rows=1, cols=1).cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]]
        for paragraph in places:
            run = paragraph.add_run('中文123%（测试）-')
            builder._set_run_font(run, size=14, bold=True, font_name='宋体', western_font='Calibri')
            run._r.rPr.rFonts.set(qn('w:asciiTheme'), 'minorHAnsi')
            run._r.rPr.rFonts.set(qn('w:hAnsiTheme'), 'minorHAnsi')
        doc.styles['Normal'].font.name = 'Calibri'
        doc.styles['Normal'].font.size = Pt(16)
        doc.styles['Normal'].element.rPr.rFonts.set(qn('w:eastAsia'), '仿宋_GB2312')
        builder._finalize_western_fonts(doc)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'final.docx'
            doc.save(path)
            saved = Document(path)
            for part in saved.part.package.parts:
                if not str(part.partname).startswith(('/word/document', '/word/header', '/word/footer', '/word/styles')):
                    continue
                if not hasattr(part, 'element'):
                    continue
                for fonts in part.element.iter(qn('w:rFonts')):
                    for slot in ('ascii', 'hAnsi', 'cs'):
                        self.assertEqual(fonts.get(qn('w:' + slot)), 'Times New Roman')
                    self.assertIsNone(fonts.get(qn('w:asciiTheme')))
                    self.assertIsNone(fonts.get(qn('w:hAnsiTheme')))
            for paragraph in places:
                self.assertEqual(paragraph.text, '中文123%（测试）-')
                self.assert_font(paragraph.runs[0], '宋体', 14)
                self.assertTrue(paragraph.runs[0].bold)

    def test_saved_format_in_every_report_skill(self):
        content = {
            'cover': {'title_lines': ['立项报告（试行2026）']},
            'blocks': [
                {'type': 'h1', 'text': '一、公司情况'},
                {'type': 'h2', 'text': '（一）基本信息'},
                {'type': 'h3', 'text': '1.经营情况'},
                {'type': 'h4', 'text': '（1）业务说明（量产2026）'},
                {'type': 'p', 'text': '正文ABC（口径(123)说明）结束50%', 'bold': True},
                {'type': 'bullet', 'items': ['经营资质（有效期2026）']},
                {'type': 'tnote', 'text': '单位（万元）'},
                {'type': 'table', 'header': ['项目（2026）'], 'rows': [['数值(ABC 123%)']]},
            ],
            'attachments': ['附件一：《实施方案（试行）》；', '附件2.测算表！'],
        }
        for skill in SKILLS:
            with self.subTest(skill=skill), tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp)
                source, output = path / 'input.json', path / 'output.docx'
                source.write_text(json.dumps(content, ensure_ascii=False), encoding='utf-8')
                result = subprocess.run([sys.executable, '-X', 'utf8',
                                         str(REPO / skill / 'scripts/build_docx.py'),
                                         str(source), str(output)], capture_output=True)
                self.assertEqual(result.returncode, 0, result.stderr.decode('utf-8', errors='replace'))
                doc = Document(output)
                table_paragraphs = [p for table in doc.tables for row in table.rows
                                    for cell in row.cells for p in cell.paragraphs]
                paragraphs = doc.paragraphs + table_paragraphs
                expected_spans = {'（试行2026）', '（一）', '（1）', '（口径(123)说明）',
                                  '（有效期2026）', '（万元）', '（2026）', '(ABC 123%)',
                                  '（试行）', '（量产2026）'}
                found = set()
                for paragraph in paragraphs:
                    for run in paragraph.runs:
                        if run.text.startswith(('（', '(')):
                            found.add(run.text)
                            font = '仿宋_GB2312' if run.text == '（1）' else '楷体_GB2312'
                            self.assert_font(run, font, 10.5 if paragraph in table_paragraphs else 16)
                self.assertEqual(found, expected_spans)
                h4 = next(p for p in doc.paragraphs if p.text.startswith('（1）'))
                self.assertTrue(all(run.bold for run in h4.runs))
                self.assert_font(h4.runs[1], '仿宋_GB2312', 16)
                for row in doc.tables[0].rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                if run.text:
                                    self.assertEqual(run.font.size.pt, 10.5)
                body = next(p for p in doc.paragraphs if p.text.startswith('正文ABC'))
                self.assertTrue(all(run.bold for run in body.runs))
                self.assertEqual(body.runs[0]._r.rPr.rFonts.get(qn('w:ascii')), 'Times New Roman')
                self.assertEqual(body.runs[0]._r.rPr.rFonts.get(qn('w:eastAsia')), '仿宋_GB2312')
                for paragraph in doc.paragraphs:
                    if not paragraph.text or paragraph.text.startswith('单位'):
                        continue
                    spacing = 30 if paragraph.text.startswith('立项报告') else 28
                    self.assertEqual(paragraph.paragraph_format.line_spacing_rule, WD_LINE_SPACING.EXACTLY)
                    self.assertEqual(paragraph.paragraph_format.line_spacing.pt, spacing)
                self.assertEqual([p.text for p in doc.paragraphs if p.text.startswith('附件')],
                                 ['附件1.实施方案（试行）', '附件2.测算表'])
                self.assertTrue(doc.settings.odd_and_even_pages_header_footer)
                self.assertNotEqual(doc.sections[0].footer.part.partname,
                                    doc.sections[0].even_page_footer.part.partname)
                for footer, align in ((doc.sections[0].footer, WD_ALIGN_PARAGRAPH.RIGHT),
                                      (doc.sections[0].even_page_footer, WD_ALIGN_PARAGRAPH.LEFT)):
                    paragraph = footer.paragraphs[0]
                    self.assertEqual(paragraph.alignment, align)
                    self.assertEqual(paragraph.text, '-1-')
                    self.assertEqual(len(paragraph.runs), 3)
                    for run in paragraph.runs:
                        self.assert_font(run, '宋体', 14)
                    self.assertEqual([n.text for n in paragraph._p.xpath('.//w:instrText')], ['PAGE'])

    def test_text_helpers_and_single_attachment(self):
        spec = importlib.util.spec_from_file_location('helpers', REPO / SKILLS[0] / 'scripts/docx_format_helpers.py')
        helpers = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(helpers)
        self.assertEqual(helpers.attachment_lines(['附件1.《工作方案（试行）》。', '《》']),
                         ['附件：工作方案（试行）'])
        self.assertEqual(helpers.normalize_attachment_name('2026年度计划'), '2026年度计划')
        self.assertEqual(helpers.normalize_attachment_name('附件2026年度说明'), '附件2026年度说明')
        self.assertEqual(helpers.parenthesized_spans('正文（未闭合'), [])
        self.assertEqual(helpers.parenthesized_spans('a（b(c)d）e'), [(1, 8)])


if __name__ == '__main__':
    unittest.main()
