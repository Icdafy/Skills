#!/usr/bin/env python3
"""跨技能公文格式一致性回归测试：以 yiti-skill 为基准。

同一份内容分别交给 yiti-skill、officialese-skill 和立项报告三技能的生成器，逐项
核对共同的公文格式标准（yiti-skill/references/format-rules.md）：标题与正文的括号
楷体及字号、四级标题序号、标题/正文行距、表格、附件说明、页脚与全角转半角。
soe-post-investment-report 的格式由其自带校验器与 tests/test_skill.py 覆盖。

各生成器以子进程调用 CLI，避免同名模块在同一进程内串味。
"""

from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn

REPO = Path(__file__).resolve().parents[2]

TITLE = "关于测试事项（2026年）的报告"
H1, H2, H3 = "一、总体情况", "（一）基本信息", "1.经营情况"
H4 = "（1）业务说明（量产２０２６）"
BODY = "正文ＡＢＣ（口径说明）增长１２％"
TABLE = [["项目（单位）", "数值"], ["收入（万元）", "1,234.56"]]
ATTACHMENTS = ["《实施方案（试行）》", "测算表"]


def _yiti(tmp: Path, out: Path) -> list[str]:
    spec = tmp / "yiti.json"
    spec.write_text(json.dumps({
        "kind": "exit", "title": TITLE, "intro": [BODY],
        "blocks": [{"type": "h1", "text": H1}, {"type": "h2", "text": H2},
                   {"type": "h3", "text": H3}, {"type": "h4", "text": H4},
                   {"type": "table", "header": TABLE[0], "rows": TABLE[1:]}],
        "attachments": ATTACHMENTS,
        "signature": {"issuer": "某某投资管理有限公司", "date": "2026年9月1日"},
    }, ensure_ascii=False), encoding="utf-8")
    return [sys.executable, str(REPO / "yiti-skill/scripts/create_yiti_docx.py"), str(spec), str(out)]


def _officialese(tmp: Path, out: Path) -> list[str]:
    spec = tmp / "official.json"
    spec.write_text(json.dumps({
        "title": TITLE, "body": [BODY],
        "sections": [{"heading": H1, "level": 1, "children": [
            {"heading": H2, "level": 2, "children": [
                {"heading": H3, "level": 3, "children": [
                    {"heading": H4, "level": 4, "paragraphs": [{"table": TABLE}]}]}]}]}],
        "attachments": ATTACHMENTS,
        "issuer": "某某投资管理有限公司", "date": "2026年9月1日",
    }, ensure_ascii=False), encoding="utf-8")
    return [sys.executable, str(REPO / "officialese-skill/scripts/create_official_docx.py"),
            str(out), "--input", str(spec)]


def _trio(skill: str):
    def command(tmp: Path, out: Path) -> list[str]:
        spec = tmp / "content.json"
        spec.write_text(json.dumps({
            "cover": {"title_lines": [TITLE]},
            "blocks": [{"type": "p", "text": BODY},
                       {"type": "h1", "text": H1}, {"type": "h2", "text": H2},
                       {"type": "h3", "text": H3}, {"type": "h4", "text": H4},
                       {"type": "table", "header": TABLE[0], "rows": TABLE[1:]}],
            "attachments": ATTACHMENTS,
            "signature": {"issuer": "某某投资管理有限公司", "date": "2026年9月1日"},
        }, ensure_ascii=False), encoding="utf-8")
        return [sys.executable, str(REPO / skill / "scripts/build_docx.py"), str(spec), str(out)]
    return command


GENERATORS = {
    "yiti-skill": _yiti,
    "officialese-skill": _officialese,
    "gongsi-qingkuang": _trio("gongsi-qingkuang"),
    "hangye-fenxi": _trio("hangye-fenxi"),
    "zhuying-yewu-fenxi": _trio("zhuying-yewu-fenxi"),
}


def east_asia(run) -> str:
    return run._r.rPr.rFonts.get(qn("w:eastAsia"))


class UnifiedFormatTests(unittest.TestCase):
    def build(self, skill: str) -> Document:
        tmp = Path(self._tmp.name) / skill
        tmp.mkdir()
        out = tmp / "out.docx"
        proc = subprocess.run(GENERATORS[skill](tmp, out), capture_output=True, text=True,
                              encoding="utf-8", errors="replace")
        self.assertEqual(proc.returncode, 0, proc.stderr)
        return Document(out)

    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def test_every_generator_follows_the_yiti_standard(self) -> None:
        for skill in GENERATORS:
            with self.subTest(skill=skill):
                self.check(self.build(skill))

    def check(self, doc: Document) -> None:
        paragraphs = doc.paragraphs
        by_text = {p.text: p for p in paragraphs if p.text}

        # 主标题：二号方正小标宋；标题内括号楷体_GB2312，与标题同为二号。
        title = by_text[TITLE]
        self.assertEqual(title.paragraph_format.line_spacing.pt, 30)
        for run in title.runs:
            if run.text.startswith("（"):
                self.assertEqual((east_asia(run), run.font.size.pt), ("楷体_GB2312", 22))
            elif run.text:
                self.assertEqual((east_asia(run), run.font.size.pt), ("方正小标宋简体", 22))

        # 正文：三号仿宋、固定 28 磅；括号三号楷体；全角数字、字母和％已转半角。
        body = by_text["正文ABC（口径说明）增长12%"]
        self.assertEqual(body.paragraph_format.line_spacing.pt, 28)
        self.assertEqual(body.paragraph_format.line_spacing_rule, WD_LINE_SPACING.EXACTLY)
        paren = next(run for run in body.runs if run.text == "（口径说明）")
        self.assertEqual((east_asia(paren), paren.font.size.pt), ("楷体_GB2312", 16))

        # 各级标题：固定 30 磅、首行缩进两字；字体与加粗按层级。
        expected = {H1: ("黑体", False), H2: ("楷体_GB2312", True), H3: ("仿宋_GB2312", True)}
        for text, (font, bold) in expected.items():
            heading = by_text[text]
            self.assertEqual(heading.paragraph_format.line_spacing.pt, 30, text)
            for run in heading.runs:
                if run.text:
                    self.assertEqual((east_asia(run), bool(run.bold)), (font, bold), text)
        h4 = by_text["（1）业务说明（量产2026）"]
        self.assertEqual(h4.paragraph_format.line_spacing.pt, 30)
        self.assertTrue(all(run.bold for run in h4.runs if run.text))
        serial = h4.runs[0]
        self.assertTrue(serial.text.startswith("（1）"))
        self.assertEqual(east_asia(serial), "仿宋_GB2312")  # 序号随标题，不改楷体
        note = next(run for run in h4.runs if run.text == "（量产2026）")
        self.assertEqual((east_asia(note), note.font.size.pt), ("楷体_GB2312", 16))

        # 表格：五号、单倍行距、居中；表头加粗且无底纹、跨页重复；行不跨页。
        table = doc.tables[0]
        for r_idx, row in enumerate(table.rows):
            tr_pr = row._tr.trPr
            self.assertIsNotNone(tr_pr.find(qn("w:cantSplit")))
            if r_idx == 0:
                self.assertIsNotNone(tr_pr.find(qn("w:tblHeader")))
            for cell in row.cells:
                tc_pr = cell._tc.tcPr
                self.assertIsNone(tc_pr.find(qn("w:shd")) if tc_pr is not None else None)
                for paragraph in cell.paragraphs:
                    rule = paragraph.paragraph_format.line_spacing_rule
                    if rule is None:  # 未显式设置时继承 Normal，须为单倍行距
                        self.assertIsNone(doc.styles["Normal"].paragraph_format.line_spacing)
                    else:
                        self.assertEqual(rule, WD_LINE_SPACING.SINGLE)
                    for run in paragraph.runs:
                        if not run.text:
                            continue
                        self.assertEqual(run.font.size.pt, 10.5)
                        self.assertEqual(bool(run.bold), r_idx == 0)
                        font = "楷体_GB2312" if run.text.startswith("（") else "仿宋_GB2312"
                        self.assertEqual(east_asia(run), font)

        # 附件说明：附件：1.XXX / 2.XXX，"2."与"1."对齐，名称无书名号。
        texts = [p.text for p in paragraphs]
        first = texts.index("附件：1.实施方案（试行）")
        self.assertEqual(texts[first - 1], "")
        self.assertEqual(texts[first + 1], "2.测算表")
        starts = []
        for p in paragraphs[first:first + 2]:
            fmt = p.paragraph_format
            starts.append(fmt.left_indent.pt + fmt.first_line_indent.pt)
        self.assertAlmostEqual(starts[0], 32, places=1)
        self.assertAlmostEqual(starts[1], 80, places=1)

        # 落款：署名右空四字，日期在署名下居中。
        issuer = by_text["某某投资管理有限公司"]
        date = by_text["2026年9月1日"]
        self.assertAlmostEqual(issuer.paragraph_format.right_indent.pt, 64, places=1)
        self.assertGreater(date.paragraph_format.right_indent.pt, 64)

        # 全文（不含页脚）西文字体 Times New Roman。
        for run in doc.element.body.iter(qn("w:r")):
            fonts = run.find(qn("w:rPr") + "/" + qn("w:rFonts"))
            if fonts is not None:
                self.assertEqual(fonts.get(qn("w:ascii")), "Times New Roman")

        # 页脚：奇偶页不同，完整 -1- 四号宋体，四个字体槽均为宋体。
        self.assertTrue(doc.settings.odd_and_even_pages_header_footer)
        section = doc.sections[0]
        for footer in (section.footer, section.even_page_footer):
            paragraph = footer.paragraphs[0]
            self.assertEqual(paragraph.text, "-1-")
            for run in paragraph.runs:
                for slot in ("ascii", "hAnsi", "cs", "eastAsia"):
                    self.assertEqual(run._r.rPr.rFonts.get(qn("w:" + slot)), "宋体")
                self.assertEqual(run.font.size.pt, 14)


if __name__ == "__main__":
    unittest.main()
